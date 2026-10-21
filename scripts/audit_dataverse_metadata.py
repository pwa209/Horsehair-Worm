"""Native metadata inventory only: preserve source rows; never assign treatments.

Run on an allocated Narval worker with the complete official search response as
--roster. The output is a fresh immutable investigation case, not a study-ready
crosswalk. Trajectories and binary tracking files are never downloaded.
"""
from __future__ import annotations
import argparse
import concurrent.futures
import csv
from datetime import date, datetime, time, timezone
import hashlib
import io
import json
import os
import re
from pathlib import Path
import socket
import time as timer
import urllib.request
from openpyxl import load_workbook

LIMIT = 200_000
DATASET = 'doi:10.7910/DVN/LTMCFR'
END_TIME = float('inf')
RUN_ID = 'uninitialized'
DIRECTORY_METADATA = {}

def digest(data: bytes, kind: str = 'sha256') -> str:
    return hashlib.new(kind, data).hexdigest()

def native_value(value):
    if isinstance(value, (datetime, date, time)):
        return {'type': type(value).__name__, 'value': value.isoformat()}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return {'type': type(value).__name__, 'value': str(value)}

def request(url: str, limit: int = LIMIT) -> bytes:
    if timer.monotonic() >= END_TIME:
        raise TimeoutError('Total native audit deadline reached')
    req = urllib.request.Request(url, headers={'User-Agent': 'Horsehair-study-metadata-audit/1.0'})
    with urllib.request.urlopen(req, timeout=35) as response:
        if response.status != 200:
            raise ValueError(f'Unexpected HTTP status {response.status}')
        data = response.read(limit + 1)
    if len(data) > limit:
        raise ValueError('Metadata-only byte bound exceeded')
    return data

def inspect_workbook(data: bytes) -> list[dict]:
    parsed = []
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        for sheet in workbook:
            if sheet.max_row > 100_000 or sheet.max_column > 128:
                raise ValueError('Unexpected original metadata dimensions')
            rows = [[native_value(cell.value) for cell in row] for row in sheet.iter_rows()]
            fields = []
            # Report explicitly written header candidates only. Their biological
            # meaning, units, row links and completeness require a later review.
            for ri, row in enumerate(rows[:10]):
                for ci, val in enumerate(row):
                    if isinstance(val, str) and re.search(r'(?i)(expos|infect|treat|\brx\b|time|clock|animal|\broi\b|reagent|stock|control|geno|well|donor|recipient|family)', val):
                        fields.append({'excel_row': ri+1, 'excel_column': ci+1, 'literal': val})
            parsed.append({'sheet': sheet.title, 'rows': rows, 'explicit_header_candidates': fields,
                           'interpretation': 'Original cells with Excel row 1 at rows[0]; formulas retained, no scientific values imputed'})
    finally:
        workbook.close()
    return parsed

def checkpoint(item: dict, case: Path) -> dict | None:
    path = case / 'downloads' / str(item['file_id']) / 'verified.json'
    if not path.is_file():
        return None
    record = json.loads(path.read_bytes())
    folder = Path(record['verified_artifact_directory'])
    if not folder.resolve().is_relative_to(path.parent.resolve()):
        raise ValueError('Checkpoint artifact escaped its original datafile directory')
    for filename, key in [('original.xlsx', 'workbook_sha256'), ('provider-file-metadata.json', 'file_metadata_sha256'), ('original-cells.json', 'cells_sha256')]:
        content = (folder / filename).read_bytes()
        if digest(content) != record[key]:
            raise ValueError('Verified metadata checkpoint changed')
    data = (folder / 'original.xlsx').read_bytes()
    if record['file_id'] != str(item['file_id']) or len(data) != item['size_in_bytes'] or digest(data, 'md5') != item['checksum']['value']:
        raise ValueError('Checkpoint no longer matches original provider metadata')
    return record

def version_directory_map(content: bytes, rows: list[dict], expected_files: int = 6068) -> dict:
    obj = json.loads(content)
    version = obj['data']
    if (obj['status'] != 'OK' or version['versionNumber'] != 1 or version['versionMinorNumber'] != 0
        or len(version['files']) != expected_files):
        raise ValueError('Dataset version disagrees with complete original V1.0 source')
    original = {str(r['dataFile']['id']): r for r in version['files']}
    if len(original) != len(version['files']):
        raise ValueError('Dataset version duplicates original file identities')
    for item in rows:
        entry = original[str(item['file_id'])]
        if (entry['dataFile']['filesize'] != item['size_in_bytes'] or entry['dataFile']['checksum'] != item['checksum']
            or entry.get('label') != item['name'] or entry.get('restricted') is not False
            or not isinstance(entry.get('directoryLabel'), str) or not entry['directoryLabel']):
            raise ValueError('Dataset version file binding disagrees with provider roster')
    return {str(item['file_id']): json.dumps(original[str(item['file_id'])]).encode('utf-8') for item in rows}

def one(item: dict, case: Path) -> dict:
    fid = str(item['file_id'])
    result = {'file_id': fid, 'filename': item['name'], 'provider_bytes': item['size_in_bytes'],
              'provider_md5': item['checksum']['value'], 'source_url': item['url'], 'attempts': []}
    try:
        existing = checkpoint(item, case)
        if existing is not None:
            return dict(existing, reused_verified_checkpoint=True)
    except Exception as error:
        return dict(result, status='changed_verified_checkpoint', error=str(error))
    for attempt in range(1, 4):
        try:
            data = request(f'https://dataverse.harvard.edu/api/access/datafile/{fid}')
            parent = case / 'downloads' / fid
            parent.mkdir(exist_ok=True)
            folder = parent / (RUN_ID + '-attempt-' + str(attempt))
            folder.mkdir()
            (folder / 'original-candidate.xlsx').write_bytes(data)
            if len(data) != item['size_in_bytes'] or digest(data, 'md5') != item['checksum']['value']:
                raise ValueError('Original workbook disagrees with provider bytes or MD5')
            metadata = DIRECTORY_METADATA.get(fid)
            if metadata is None:
                metadata = request(f'https://dataverse.harvard.edu/api/files/{fid}/metadata')
            meta = json.loads(metadata)
            if meta.get('label') != item['name'] or meta.get('restricted') is not False:
                raise ValueError('Original file metadata label or access disagrees with roster')
            if not isinstance(meta.get('directoryLabel'), str) or not meta['directoryLabel']:
                raise ValueError('Original experiment directory is absent')
            parsed = inspect_workbook(data)
            cells = json.dumps(parsed, indent=2, allow_nan=False).encode('utf-8')
            (folder / 'original.xlsx').write_bytes(data)
            (folder / 'provider-file-metadata.json').write_bytes(metadata)
            (folder / 'original-cells.json').write_bytes(cells)
            result.update(status='verified_original_metadata', directory_label=meta['directoryLabel'],
                          verified_artifact_directory=str(folder),
                          workbook_sha256=digest(data), file_metadata_sha256=digest(metadata), cells_sha256=digest(cells),
                          sheets=[{'name': s['sheet'], 'rows': len(s['rows']),
                                   'explicit_header_candidates': s['explicit_header_candidates'],
                                   'first_nonempty_rows': [{'excel_row': i+1, 'values': r} for i,r in enumerate(s['rows']) if any(v is not None for v in r)][:3]}
                                  for s in parsed])
            pending = parent / (RUN_ID + '-verified.partial')
            pending.write_text(json.dumps(result, indent=2, allow_nan=False), encoding='utf-8')
            os.replace(pending, parent / 'verified.json')
            return result
        except Exception as error:
            result['attempts'].append({'attempt': attempt, 'error': type(error).__name__ + ': ' + str(error),
                                       'utc': datetime.now(timezone.utc).isoformat()})
            if timer.monotonic() >= END_TIME:
                break
            if attempt < 3:
                timer.sleep(attempt)
    result['status'] = 'unresolved_retrieval_or_parser'
    return result

def main():
    global END_TIME, RUN_ID, DIRECTORY_METADATA
    parser = argparse.ArgumentParser()
    parser.add_argument('--roster', required=True, type=Path)
    parser.add_argument('--roster-sha256', required=True)
    parser.add_argument('--case', required=True, type=Path)
    parser.add_argument('--max-seconds', type=int, default=6300)
    parser.add_argument('--resume', action='store_true')
    parser.add_argument('--primary-summary', type=Path)
    parser.add_argument('--primary-summary-sha256')
    parser.add_argument('--framework-release', required=True)
    parser.add_argument('--framework-source-sha256', required=True)
    args = parser.parse_args()
    if not os.environ.get('SLURM_JOB_ID') or os.environ.get('SLURM_CLUSTER_NAME', '').lower() != 'narval' or os.getuid() != 3164768:
        raise ValueError('An allocated Narval worker owned by pwa209 is required')
    if not 60 <= args.max_seconds <= 6500:
        raise ValueError('Native audit deadline must fit a two-hour allocation')
    END_TIME = timer.monotonic() + args.max_seconds
    RUN_ID = 'job-' + str(os.environ['SLURM_JOB_ID'])
    framework = Path(args.framework_release).resolve()
    from horsehair.core import source_identity
    import horsehair.core
    if not framework.is_absolute() or not Path(horsehair.core.__file__).resolve().is_relative_to(framework / 'src'):
        raise ValueError('Runtime framework must be the explicit immutable qualified release')
    if source_identity(framework)['sha256'] != args.framework_source_sha256:
        raise ValueError('Runtime framework source disagrees with qualified release binding')
    study_root = Path('/scratch/pwa209/horsehair-worm').resolve()
    destination = args.case.resolve()
    if not destination.is_relative_to(study_root / 'maintenance') or destination == study_root / 'maintenance':
        raise ValueError('Output case must be a fresh personal-study maintenance subdirectory')
    roster_bytes = args.roster.read_bytes()
    if digest(roster_bytes) != args.roster_sha256:
        raise ValueError('Official metadata roster changed')
    search = json.loads(roster_bytes)
    payload = search['data']
    rows = payload['items']
    if search['status'] != 'OK' or payload['total_count'] != 875 or len(rows) != 875:
        raise ValueError('Full inspected 875-file official roster is required')
    if len({r['file_id'] for r in rows}) != len(rows):
        raise ValueError('Duplicate original datafile identity')
    for row in rows:
        if (row['dataset_persistent_id'] != DATASET or row['dataset_id'] != '6802078'
            or row['file_content_type'] != 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            or row['restricted'] or not str(row['file_id']).isdigit()
            or row['checksum']['type'] != 'MD5' or row['size_in_bytes'] > LIMIT):
            raise ValueError('Roster contains an unrelated or non-metadata original file')
    if args.case.exists():
        if not args.resume or (args.case / 'official-roster.json').read_bytes() != roster_bytes:
            raise ValueError('Existing investigation requires explicit resume and identical original roster')
    else:
        args.case.mkdir()
        (args.case / 'downloads').mkdir()
        (args.case / 'runs').mkdir()
        (args.case / 'official-roster.json').write_bytes(roster_bytes)
    run_case = args.case / 'runs' / RUN_ID
    run_case.mkdir()
    primary = []
    if args.primary_summary:
        content = args.primary_summary.read_bytes()
        if digest(content) != args.primary_summary_sha256:
            raise ValueError('Original pinned classifier summary changed')
        primary = [r['experiment_name'] for r in csv.DictReader(io.StringIO(content.decode('utf-8-sig')))]
    intent = {'started_utc': datetime.now(timezone.utc).isoformat(), 'purpose': 'original-source metadata inventory only',
              'scientific_readiness': False, 'source_dataset': DATASET,
              'source_roster_sha256': args.roster_sha256, 'script_sha256': digest(Path(__file__).read_bytes()),
              'framework_release': args.framework_release, 'framework_source_sha256': args.framework_source_sha256,
              'primary_summary_sha256': args.primary_summary_sha256,
              'native': {'job_id': os.environ['SLURM_JOB_ID'], 'host': socket.gethostname(), 'user_id': os.getuid()},
              'expected_files': len(rows), 'expected_workbook_bytes': sum(r['size_in_bytes'] for r in rows)}
    (run_case / 'intent.json').write_text(json.dumps(intent, indent=2), encoding='utf-8')
    version_attempts = []
    for version_url in ['https://dataverse.harvard.edu/api/datasets/6802078/versions/1.0',
                        'https://dataverse.harvard.edu/api/datasets/:persistentId/versions/1.0?persistentId=doi:10.7910/DVN/LTMCFR']:
        try:
            content = request(version_url, limit=20_000_000)
            (run_case / 'dataset-version-metadata.json').write_bytes(content)
            DIRECTORY_METADATA = version_directory_map(content, rows)
            version_attempts.append({'url': version_url, 'status': 'verified_dataset_version', 'sha256': digest(content), 'files': 6068})
            break
        except Exception as error:
            version_attempts.append({'url': version_url, 'status': 'unavailable_or_mismatched', 'error': str(error)})
    (run_case / 'dataset-version-attempts.json').write_text(json.dumps(version_attempts, indent=2), encoding='utf-8')
    results = []
    with (run_case / 'inventory-progress.jsonl').open('x', encoding='utf-8') as log:
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            jobs = [pool.submit(one, row, args.case) for row in rows]
            for future in concurrent.futures.as_completed(jobs):
                result = future.result()
                results.append(result)
                log.write(json.dumps(result, allow_nan=False) + '\n')
                log.flush()
                if len(results) % 25 == 0:
                    print(json.dumps({'observed': len(results), 'expected': len(rows),
                                      'unresolved': sum(r['status'] != 'verified_original_metadata' for r in results)}), flush=True)
    results.sort(key=lambda r: int(r['file_id']))
    unresolved = [r['file_id'] for r in results if r['status'] != 'verified_original_metadata']
    primary_matches = {name: [r['file_id'] for r in results if r.get('directory_label', '').split('/')[-1] == name] for name in primary}
    summary = dict(intent, finished_utc=datetime.now(timezone.utc).isoformat(),
                   status='complete_source_inventory' if not unresolved else 'partial_source_inventory',
                   verified_files=len(results)-len(unresolved), unresolved_file_ids=unresolved,
                   inventories=results,
                   primary_source_exact_directory_matches=primary_matches,
                   missing_or_ambiguous_primary_sources=[name for name, matches in primary_matches.items() if len(matches) != 1],
                   prohibited_inferences=['No exposure assignment from survival, filenames, Day labels, or blank Rx',
                                          'No cross-session animal identity or shared-control reuse inferred from board/ROI coincidence',
                                          'No gene/reagent or family assignment from original source conflicts',
                                          'No source crosswalk or H1/H2/H3 readiness cleared by inventory completion'])
    encoded = json.dumps(summary, indent=2, allow_nan=False).encode('utf-8')
    (run_case / 'inventory.json').write_bytes(encoded)
    print(json.dumps({'status': summary['status'], 'verified_files': summary['verified_files'],
                      'unresolved_file_ids': unresolved, 'inventory_sha256': digest(encoded)}), flush=True)
    if unresolved:
        raise SystemExit(2)

if __name__ == '__main__':
    main()
