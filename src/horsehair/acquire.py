"""Versioned public-data acquisition; source documents are data, never commands.

The discovery command transfers repository metadata only. Large FASTQs are gated
by a reconciled biological crosswalk. Downloads are resumable and atomically
promoted only after their declared size and provider digest have been checked.
"""
from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import io
import json
import os
import posixpath
import re
import shutil
import stat
import tarfile
import time
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import quote, urljoin, urlparse

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

CATALOG = {
    "fly": {"doi": "10.5281/zenodo.7464925", "record": "7464925", "expected_file": "Summit paper data.zip", "expected_md5": "f7c1335d7d831150dbf45b511173925e", "companion": "https://lab.debivort.org/zombie-summiting/"},
    "polarization": {"doi": "10.5061/dryad.jh9w0vtkb"},
    "mating": {"doi": "10.5061/dryad.4j0zpc8jw"},
    "mantis_rna": {"accession": "GSE237182", "bioproject": "PRJNA994131"},
    "ant_rna": {"accession": "PRJNA918668", "primary_year": "2021", "expected_brains": 60},
    "ant_reference": {"accession": "GCA_003651465.1", "organism": "Formica exsecta"},
    "fly_reference": {"release": "FB2026_03", "genome_release": "r6.69", "organism": "Drosophila melanogaster", "license_url": "https://wiki.flybase.org/wiki/FlyBase:About"},
}
ENA_FIELDS = ["run_accession", "experiment_accession", "sample_accession", "secondary_sample_accession", "sample_alias", "sample_title", "scientific_name", "experiment_title", "library_name", "library_layout", "read_count", "base_count", "fastq_ftp", "fastq_md5", "fastq_bytes"]


class AcquisitionError(RuntimeError):
    pass


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def canonical_hash(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()


def file_hash(path: str | Path, algorithm: str = "sha256") -> str:
    digest = hashlib.new(algorithm.replace("-", ""))
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_digests(path: str | Path, algorithms: set[str]) -> dict[str, str]:
    contexts = {algorithm.replace("-", ""): hashlib.new(algorithm.replace("-", "")) for algorithm in algorithms}
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            for context in contexts.values():
                context.update(chunk)
    return {name: context.hexdigest() for name, context in contexts.items()}


def atomic_json(path: str | Path, value: Any) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(value, handle, indent=2, sort_keys=True, ensure_ascii=False)
        handle.write("\n")
        handle.flush()
        os.fsync(handle.fileno())
    temporary.replace(path)


def http_session() -> requests.Session:
    session = requests.Session()
    session.headers.update({"User-Agent": "HorsehairResearch/1.0 public-metadata-acquisition", "Accept-Encoding": "identity"})
    retry = Retry(total=4, backoff_factor=1.0, status_forcelist=[429, 500, 502, 503, 504], allowed_methods=["GET", "HEAD"])
    session.mount("https://", HTTPAdapter(max_retries=retry))
    return session


def checked_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise AcquisitionError(f"Only credential-free HTTPS source URLs are permitted: {url!r}")
    if parsed.hostname in {"localhost", "127.0.0.1", "::1"}:
        raise AcquisitionError("Localhost is not a public data source")
    return url


def safe_destination(root: str | Path, relative: str) -> Path:
    root = Path(root).resolve()
    normalized = relative.replace("\\", "/")
    pure = PurePosixPath(normalized)
    if pure.is_absolute() or re.match(r"^[A-Za-z]:", normalized) or ".." in pure.parts or not pure.parts or ":" in normalized:
        raise AcquisitionError(f"Unsafe relative destination: {relative!r}")
    destination = (root / Path(*pure.parts)).resolve()
    if not destination.is_relative_to(root):
        raise AcquisitionError(f"Destination escapes root: {relative!r}")
    return destination


def metadata_get(session: requests.Session, url: str) -> dict:
    response = session.get(checked_url(url), timeout=(20, 90))
    response.raise_for_status()
    try:
        return response.json()
    except ValueError as error:
        raise AcquisitionError(f"Repository returned non-JSON metadata at {url}") from error


def asset(*, source: str, name: str, url: str, version: str, rights: str, size: int | None = None, checksum: str | None = None, kind: str = "data", **extra: Any) -> dict:
    return {"asset_id": f"{source}:{name}", "source_id": source, "relative_path": f"raw/{source}/{name}", "url": checked_url(url), "version": str(version), "bytes": size, "provider_checksum": checksum, "rights": rights, "kind": kind, **extra}


def discover_zenodo(session: requests.Session) -> tuple[dict, list[dict]]:
    item = CATALOG["fly"]
    metadata = metadata_get(session, f"https://zenodo.org/api/records/{item['record']}")
    if str(metadata.get("id")) != item["record"] or metadata.get("doi") != item["doi"]:
        raise AcquisitionError("Zenodo record identity changed")
    files = metadata.get("files", [])
    matches = [f for f in files if f["key"] == item["expected_file"]]
    if len(matches) != 1 or matches[0].get("checksum") != "md5:" + item["expected_md5"]:
        raise AcquisitionError("Publication-associated fly archive missing or provider digest changed; create an explicit amendment")
    rights = metadata["metadata"].get("license", {}).get("id", "unspecified")
    assets = [asset(source="fly", name=f["key"], url=f["links"]["self"], version=item["record"], rights=rights, size=int(f["size"]), checksum=f["checksum"], kind="archive", doi=item["doi"]) for f in files]
    return metadata, assets


def discover_dryad(session: requests.Session, source: str) -> tuple[dict, list[dict]]:
    doi = CATALOG[source]["doi"]
    base = "https://datadryad.org"
    metadata = metadata_get(session, base + "/api/v2/datasets/" + quote("doi:" + doi, safe=""))
    if metadata.get("identifier") != "doi:" + doi or metadata.get("curationStatus") != "Published":
        raise AcquisitionError(f"Dryad {source} is not the requested published dataset")
    version_url = urljoin(base, metadata["_links"]["stash:version"]["href"])
    version = version_url.rsplit("/", 1)[-1]
    next_url: str | None = version_url + "/files"
    files: list[dict] = []
    seen: set[str] = set()
    while next_url:
        if next_url in seen:
            raise AcquisitionError("Cyclic Dryad pagination")
        seen.add(next_url)
        listing = metadata_get(session, next_url)
        files.extend(listing.get("_embedded", {}).get("stash:files", []))
        next_link = listing.get("_links", {}).get("next", {}).get("href")
        next_url = urljoin(base, next_link) if next_link else None
    if not files:
        raise AcquisitionError(f"Dryad {source} supplied no files")
    assets = []
    for file in files:
        file_id = file["_links"]["self"]["href"].rsplit("/", 1)[-1]
        if not file_id.isdigit():
            raise AcquisitionError("Dryad metadata has an invalid public file identifier")
        # The metadata API's download link requires an API token. The public
        # browser's file-stream route is the unauthenticated data-reuse endpoint.
        public_url = base + "/downloads/file_stream/" + file_id
        assets.append(asset(source=source, name=file["path"], url=public_url, version=version, rights=metadata.get("license", "unspecified"), size=int(file["size"]), checksum=f"{file['digestType'].replace('-', '')}:{file['digest']}" if file.get("digest") else None, doi=doi, provider_api_download_url=urljoin(base, file["_links"]["stash:download"]["href"]), provider_file_id=file_id))
    return {"dataset": metadata, "files": files, "version_id": version}, assets


class _Links(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() == "a":
            self.links.extend(v for k, v in attrs if k == "href" and v is not None)


def discover_geo(session: requests.Session) -> tuple[dict, list[dict]]:
    accession = CATALOG["mantis_rna"]["accession"]
    base = f"https://ftp.ncbi.nlm.nih.gov/geo/series/GSE237nnn/{accession}/"
    response = session.get(base + "suppl/", timeout=(20, 90))
    response.raise_for_status()
    parser = _Links()
    parser.feed(response.text)
    filenames = sorted({name for name in parser.links if name.startswith(accession + "_") and "/" not in name})
    required = {f"{accession}_count_tenodera.txt.gz", f"{accession}_count_chordodes.txt.gz", f"{accession}_SuperTranscripts_tenodera.fasta.gz"}
    if not required.issubset(filenames):
        raise AcquisitionError(f"GEO processed files missing: {sorted(required - set(filenames))}")
    urls = [(name, base + "suppl/" + quote(name)) for name in filenames]
    urls.append((accession + "_family.soft.gz", base + "soft/" + accession + "_family.soft.gz"))
    assets = []
    for name, url in urls:
        head = session.head(url, allow_redirects=True, timeout=(20, 90))
        head.raise_for_status()
        size = int(head.headers["Content-Length"]) if head.headers.get("Content-Length", "").isdigit() else None
        assets.append(asset(source="mantis_rna", name=name, url=url, version="GSE237182", rights="NCBI public repository; no dataset-specific CC license asserted", size=size, accession=accession, provider_last_modified=head.headers.get("Last-Modified")))
    return {"accession": accession, "index_url": base + "suppl/", "filenames": filenames}, assets


def discover_ant_reference(session: requests.Session) -> tuple[dict, list[dict]]:
    """Freeze the source assembly plus its matching annotated RefSeq pair.

    GenBank currently distributes no GTF for this accession. The actual
    executable alignment pair is therefore the NCBI-confirmed paired RefSeq
    FASTA/GTF, explicitly recorded as an updated annotation route.
    """
    original = "GCA_003651465.1"
    paired = "GCF_003651465.1"
    report = metadata_get(session, f"https://api.ncbi.nlm.nih.gov/datasets/v2/genome/accession/{original}/dataset_report")
    records = report.get("reports", [])
    if len(records) != 1 or records[0].get("accession") != original or records[0].get("paired_accession") != paired:
        raise AcquisitionError("Publication assembly / RefSeq pairing changed")
    base = "https://ftp.ncbi.nlm.nih.gov/genomes/all/GCF/003/651/465/GCF_003651465.1_ASM365146v1/"
    sums = session.get(base + "md5checksums.txt", timeout=(20, 90))
    sums.raise_for_status()
    checksums = {}
    for line in sums.text.splitlines():
        parts = line.split(maxsplit=1)
        if len(parts) == 2:
            checksums[parts[1].lstrip("./")] = parts[0]
    names = ["GCF_003651465.1_ASM365146v1" + suffix for suffix in ["_genomic.fna.gz", "_genomic.gtf.gz", "_protein.faa.gz", "_feature_table.txt.gz", "_assembly_report.txt"]]
    assets = []
    for name in names:
        if name not in checksums:
            raise AcquisitionError(f"Reference file lacks published MD5: {name}")
        head = session.head(base + name, timeout=(20, 90))
        head.raise_for_status()
        size = int(head.headers["Content-Length"]) if head.headers.get("Content-Length", "").isdigit() else None
        assets.append(asset(source="ant_reference", name=name, url=base + name, version=paired + ":NCBI-Annotation-Release-100", rights="NCBI public reference sequence; retain repository/depositor notices", size=size, checksum="md5:" + checksums[name], kind="reference", publication_assembly=original, alignment_assembly=paired))
    return {"dataset_report": report, "checksum_url": base + "md5checksums.txt", "checksum_text": sums.text, "publication_assembly": original, "alignment_assembly": paired, "annotation": "NCBI Annotation Release 100", "method_amendment": "Published GenBank assembly has no current GTF; align and count against its matching annotated RefSeq pair, preserving this difference from the source analysis"}, assets


def discover_fly_reference(session: requests.Session) -> tuple[dict, list[dict]]:
    """Pin matching FlyBase protein/gene/GO files from one immutable release."""
    source = CATALOG["fly_reference"]
    base = "https://s3ftp.flybase.org/"
    precomputed = base + "releases/FB2026_03/precomputed_files/"
    fasta = base + "genomes/Drosophila_melanogaster/dmel_r6.69_FB2026_03/fasta/"
    checksum = session.get(fasta + "md5sum.txt", timeout=(20, 90))
    checksum.raise_for_status()
    sums = {}
    for line in checksum.text.splitlines():
        fields = line.split(maxsplit=1)
        if len(fields) == 2:
            sums[fields[1].lstrip("*./")] = fields[0]
    protein_name = "dmel-all-translation-r6.69.fasta.gz"
    if sums.get(protein_name) != "1844f571c897e51663c2eb4e144330e9":
        raise AcquisitionError("Pinned FlyBase FB2026_03 protein checksum changed; an explicit source amendment is required")
    urls = [(protein_name, fasta+protein_name, "protein_fasta"), ("fbgn_fbtr_fbpp_fb_2026_03.tsv.gz", precomputed+"genes/fbgn_fbtr_fbpp_fb_2026_03.tsv.gz", "gene_protein_map"), ("fbgn_annotation_ID_fb_2026_03.tsv.gz", precomputed+"genes/fbgn_annotation_ID_fb_2026_03.tsv.gz", "gene_catalog"), ("gene_association.fb.gz", precomputed+"go/gene_association.fb.gz", "go_annotations"), ("go-basic.obo.gz", precomputed+"ontologies/go-basic.obo.gz", "go_ontology")]
    items = []
    for name, url, role in urls:
        response = session.head(url, allow_redirects=True, timeout=(20, 90))
        response.raise_for_status()
        if not response.headers.get("Content-Length", "").isdigit():
            raise AcquisitionError(f"Pinned FlyBase asset lacks a provider byte count: {name}")
        items.append(asset(source="fly_reference", name=name, url=url, version="FB2026_03:dmel-r6.69", rights="FlyBase-produced data CC BY 4.0; retain GO and external attribution notices", size=int(response.headers["Content-Length"]), checksum="md5:"+sums[name] if name in sums else None, kind="reference", reference_role=role, provider_last_modified=response.headers.get("Last-Modified"), provider_etag=response.headers.get("ETag")))
    return {**source, "provider_md5_url": fasta+"md5sum.txt", "provider_md5_text": checksum.text, "file_roles": {item["reference_role"]: item["relative_path"] for item in items}, "method_note": "Matched fixed FB2026_03 gene/protein/GO namespaces; embedded OBO data-version must be recorded on Narval"}, items


def fetch_ena_runs(session: requests.Session, accession: str = "PRJNA918668") -> tuple[str, list[dict]]:
    response = session.get("https://www.ebi.ac.uk/ena/portal/api/filereport", params={"accession": accession, "result": "read_run", "fields": ",".join(ENA_FIELDS), "format": "tsv"}, timeout=(20, 90))
    response.raise_for_status()
    rows = list(csv.DictReader(io.StringIO(response.text), delimiter="\t"))
    if not rows or "run_accession" not in rows[0]:
        raise AcquisitionError("ENA returned an empty or invalid sequencing-run report")
    if len({row["run_accession"] for row in rows}) != len(rows):
        raise AcquisitionError("ENA run accessions are duplicated")
    return response.url, rows


def ena_fastq_assets(rows: list[dict], included_runs: set[str] | None = None, crosswalk_hash: str | None = None) -> list[dict]:
    assets = []
    for row in rows:
        run = row["run_accession"]
        addresses = row.get("fastq_ftp", "").split(";")
        digests = row.get("fastq_md5", "").split(";")
        sizes = row.get("fastq_bytes", "").split(";")
        if not addresses[0] or not (len(addresses) == len(digests) == len(sizes)):
            raise AcquisitionError(f"FASTQ address/digest/size mismatch for {run}")
        if row.get("library_layout") != "PAIRED" or len(addresses) != 2:
            raise AcquisitionError(f"Expected exactly two paired FASTQs for {run}")
        for address, digest, size in zip(addresses, digests, sizes):
            url = "https://" + address.removeprefix("ftp://")
            assets.append(asset(source="ant_rna", name=urlparse(url).path.rsplit("/", 1)[-1], url=url, version="PRJNA918668", rights="INSDC public sequencing archive; retain depositor notices", size=int(size), checksum="md5:" + digest, kind="fastq", run_accession=run, sample_accession=row["sample_accession"], eligible=included_runs is not None and run in included_runs, crosswalk_hash=crosswalk_hash))
    return assets


def discover_sources(output: str | Path, sources: list[str] | None = None) -> dict:
    output = Path(output)
    session = http_session()
    sources = sources or ["fly", "polarization", "mating", "mantis_rna", "ant_rna", "ant_reference", "fly_reference", "tool_runtime"]
    metadata: dict[str, Any] = {}
    assets: list[dict] = []
    for source in sources:
        if source == "fly":
            detail, items = discover_zenodo(session)
        elif source in {"polarization", "mating"}:
            detail, items = discover_dryad(session, source)
        elif source == "mantis_rna":
            detail, items = discover_geo(session)
        elif source == "ant_rna":
            url, rows = fetch_ena_runs(session)
            from .molecular import ant_crosswalk, fetch_ant_sample_metadata
            samples = fetch_ant_sample_metadata(rows)
            crosswalk, crosswalk_audit = ant_crosswalk(rows, samples)
            included = set(crosswalk.loc[crosswalk.included, "run_id"])
            detail = {"url": url, "rows": rows, "sample_metadata": samples, "crosswalk": crosswalk.fillna("").to_dict("records"), "crosswalk_audit": crosswalk_audit}
            items = ena_fastq_assets(rows, included, crosswalk_audit["crosswalk_hash"])
        elif source == "ant_reference":
            detail, items = discover_ant_reference(session)
        elif source == "fly_reference":
            detail, items = discover_fly_reference(session)
        elif source == "tool_runtime":
            from .molecular import tool_runtime_manifest
            runtime = tool_runtime_manifest()
            detail, items = runtime["sources"]["tool_runtime"], runtime["assets"]
        else:
            raise AcquisitionError(f"Unknown source: {source}")
        metadata[source] = detail
        assets.extend(items)
    manifest = {"schema_version": 1, "created_utc": utc_now(), "sources": metadata, "assets": assets, "status": "metadata_plan", "parent_hash": None}
    manifest["manifest_hash"] = canonical_hash(manifest)
    destination = output if output.suffix.lower() == ".json" else output / "source_manifest.json"
    if destination.exists():
        raise AcquisitionError("Manifest already exists: discovery cannot overwrite a frozen source version; use a new output directory")
    atomic_json(destination, manifest)
    return manifest


def validate_manifest(manifest: dict) -> None:
    if manifest.get("schema_version") != 1 or not manifest.get("assets"):
        raise AcquisitionError("Invalid or empty source manifest")
    expected = manifest.get("manifest_hash")
    if expected and canonical_hash({k: v for k, v in manifest.items() if k != "manifest_hash"}) != expected:
        raise AcquisitionError("Source manifest hash does not match content")
    ids, paths = set(), set()
    for item in manifest["assets"]:
        for key in ["asset_id", "source_id", "relative_path", "url", "version", "rights"]:
            if not item.get(key):
                raise AcquisitionError(f"Missing {key} in source object")
        checked_url(item["url"])
        if item["asset_id"] in ids or item["relative_path"] in paths:
            raise AcquisitionError("Duplicate asset IDs or destinations")
        ids.add(item["asset_id"])
        paths.add(item["relative_path"])
        if item.get("bytes") is not None and int(item["bytes"]) < 0:
            raise AcquisitionError("Negative declared byte count")
        digest = item.get("provider_checksum")
        if digest:
            algorithm, value = digest.split(":", 1)
            if algorithm.replace("-", "") not in {"md5", "sha256", "sha512"} or not re.fullmatch(r"[a-fA-F0-9]+", value):
                raise AcquisitionError("Invalid provider checksum")


def verify_file(path: str | Path, item: dict, *, computed_digests: dict[str, str] | None = None) -> dict:
    path = Path(path)
    size = path.stat().st_size
    if item.get("bytes") is not None and size != int(item["bytes"]):
        raise AcquisitionError(f"Byte mismatch for {path.name}: expected {item['bytes']}, received {size}")
    digest = item.get("provider_checksum")
    algorithms = {"sha256"}
    if digest:
        algorithms.add(digest.split(":", 1)[0].replace("-", ""))
    digests = computed_digests or file_digests(path, algorithms)
    if digest:
        algorithm, expected = digest.split(":", 1)
        actual = digests[algorithm.replace("-", "")]
        if actual.lower() != expected.lower():
            raise AcquisitionError(f"Provider checksum mismatch for {path.name}")
    with path.open("rb") as handle:
        header = handle.read(1024).lstrip().lower()
    if path.suffix.lower() not in {".html", ".htm"} and (header.startswith(b"<!doctype html") or header.startswith(b"<html")):
        raise AcquisitionError(f"HTML response masquerades as {path.name}")
    return {"asset_id": item["asset_id"], "path": str(path), "bytes": size, "provider_checksum": digest, "sha256": digests["sha256"], "retrieved_utc": utc_now(), "status": "verified"}


def quarantine_invalid_complete_partial(item: dict, root: str | Path, recovery_id: str) -> dict:
    """Preserve a positively identified bad complete partial before a fresh GET.

    This explicit maintenance operation never relaxes a provider checksum,
    touches a verified final file, or discards the failed bytes/provenance.
    Hashing belongs in a scheduled allocation for large source files.
    """
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,63}", recovery_id):
        raise AcquisitionError("Invalid finite recovery identity")
    if not item.get("provider_checksum") or not isinstance(item.get("bytes"), int):
        raise AcquisitionError("Recovery requires a pinned size and provider checksum")
    target = safe_destination(root, item["relative_path"])
    partial = target.with_name(target.name + ".part")
    sidecar = partial.with_name(partial.name + ".json")
    if target.exists():
        raise AcquisitionError("Final source exists; never quarantine or overwrite it")
    if partial.is_symlink() or sidecar.is_symlink() or not partial.is_file() or not sidecar.is_file():
        raise AcquisitionError("Recovery requires ordinary partial bytes and their original sidecar")
    receipt = json.loads(sidecar.read_text(encoding="utf8"))
    if receipt.get("asset_hash") != canonical_hash(item):
        raise AcquisitionError("Partial transfer provenance differs; preserve for manual source review")
    if partial.stat().st_size != item["bytes"]:
        raise AcquisitionError("Incomplete partial may resume; this recovery requires declared complete length")
    algorithm = item["provider_checksum"].split(":", 1)[0].replace("-", "")
    digests = file_digests(partial, {"sha256", algorithm})
    try:
        verify_file(partial, item, computed_digests=digests)
    except AcquisitionError as error:
        reason = str(error)
    else:
        raise AcquisitionError("Partial verifies successfully; allow ordinary verified promotion")
    # Every path is resolved inside the authorized source root before moving.
    area = safe_destination(root, str(PurePosixPath(item["relative_path"]).parent / ".quarantine" / recovery_id))
    area.mkdir(parents=True, exist_ok=False)
    record = {"status": "preserving_invalid_partial", "asset_id": item["asset_id"],
              "asset_hash": canonical_hash(item), "source_version": item["version"],
              "provider_checksum": item["provider_checksum"], "bytes": partial.stat().st_size,
              "observed_digests": digests, "reason": reason, "created_utc": utc_now(),
              "original_partial": str(partial), "original_sidecar": str(sidecar),
              "preserved_partial": str(area / partial.name), "preserved_sidecar": str(area / sidecar.name),
              "sidecar_sha256": file_hash(sidecar), "source_definition_unchanged": True}
    atomic_json(area / "recovery.json", record)
    partial.rename(area / partial.name)
    sidecar.rename(area / sidecar.name)
    record["status"] = "quarantined_invalid_partial"
    atomic_json(area / "recovery.json", record)
    return record


def download_asset(item: dict, root: str | Path, *, session: requests.Session | None = None, attempts: int = 4, bandwidth_bytes_per_second: int | None = 64 * 1024**2, max_seconds: float | None = None) -> dict:
    if bandwidth_bytes_per_second is not None and bandwidth_bytes_per_second <= 0:
        raise AcquisitionError("Transfer bandwidth ceiling must be positive")
    if max_seconds is not None and max_seconds <= 0:
        raise AcquisitionError("Transfer runtime ceiling must be positive")
    if item.get("kind") == "fastq" and (not item.get("eligible") or not item.get("crosswalk_hash")):
        raise AcquisitionError("Raw ant reads are gated until the biological crosswalk is validated")
    session = session or http_session()
    target = safe_destination(root, item["relative_path"])
    if target.exists():
        return verify_file(target, item)
    target.parent.mkdir(parents=True, exist_ok=True)
    partial = target.with_name(target.name + ".part")
    receipt_path = partial.with_name(partial.name + ".json")
    signature = canonical_hash(item)
    receipt = json.loads(receipt_path.read_text()) if receipt_path.exists() else {}
    started = time.monotonic()
    transferred = 0
    if partial.exists() and receipt.get("asset_hash") != signature:
        raise AcquisitionError(f"Partial transfer provenance differs for {target.name}; move it aside explicitly")
    for attempt in range(attempts):
        offset = partial.stat().st_size if partial.exists() else 0
        headers = {"Accept-Encoding": "identity"}
        if offset:
            headers["Range"] = f"bytes={offset}-"
            if receipt.get("etag"):
                headers["If-Range"] = receipt["etag"]
            elif receipt.get("last_modified"):
                headers["If-Range"] = receipt["last_modified"]
        try:
            with session.get(checked_url(item["url"]), headers=headers, stream=True, timeout=(20, 120)) as response:
                if response.status_code == 416 and item.get("bytes") == offset:
                    report = verify_file(partial, item)
                    partial.replace(target)
                    receipt_path.unlink(missing_ok=True)
                    report["path"] = str(target)
                    return report
                response.raise_for_status()
                if response.status_code == 206:
                    content_range = response.headers.get("Content-Range", "")
                    if not content_range.startswith(f"bytes {offset}-"):
                        raise AcquisitionError("Server returned an incorrect range")
                    mode = "ab"
                elif response.status_code == 200:
                    mode, offset = "wb", 0
                else:
                    raise AcquisitionError(f"Unexpected download status {response.status_code}")
                receipt = {"asset_hash": signature, "etag": response.headers.get("ETag"), "last_modified": response.headers.get("Last-Modified"), "url": response.url}
                atomic_json(receipt_path, receipt)
                algorithms = {"sha256"}
                if item.get("provider_checksum"):
                    algorithms.add(item["provider_checksum"].split(":",1)[0].replace("-", ""))
                contexts = {name: hashlib.new(name) for name in algorithms}
                if mode == "ab":
                    with partial.open("rb") as prefix:
                        for chunk in iter(lambda: prefix.read(8 * 1024 * 1024), b""):
                            for context in contexts.values():
                                context.update(chunk)
                with partial.open(mode) as handle:
                    for chunk in response.iter_content(chunk_size=8 * 1024 * 1024):
                        if chunk:
                            handle.write(chunk)
                            for context in contexts.values():
                                context.update(chunk)
                            transferred += len(chunk)
                            elapsed = time.monotonic() - started
                            if max_seconds is not None and elapsed >= max_seconds:
                                raise AcquisitionError("Finite transfer runtime reached; partial transfer retained for resumption")
                            if bandwidth_bytes_per_second:
                                delay = transferred / bandwidth_bytes_per_second - elapsed
                                if delay > 0:
                                    time.sleep(min(delay, 1.0))
                    handle.flush()
                    os.fsync(handle.fileno())
            report = verify_file(partial, item, computed_digests={name: context.hexdigest() for name, context in contexts.items()})
            partial.replace(target)
            receipt_path.unlink(missing_ok=True)
            report["path"] = str(target)
            return report
        except (requests.RequestException, OSError) as error:
            if attempt + 1 == attempts:
                raise AcquisitionError(f"Transfer failed for {item['asset_id']}: {error}") from error
            time.sleep(min(2 ** attempt, 8))
    raise AcquisitionError("Transfer exhausted retries")


def rna_pilot_asset_ids(manifest):
    """Fixed low/high-depth discovery brains plus required processed references."""
    source=manifest['sources']['ant_rna']
    if source.get('crosswalk_audit',{}).get('status')!='validated':raise AcquisitionError('RNA pilot requires validated biological crosswalk')
    rows=[r for r in source['crosswalk'] if str(r.get('included')).lower() in ['true','1'] and r.get('cohort')=='discovery']
    depth={}
    for row in rows:
        depth[row['biological_id']]=depth.get(row['biological_id'],0)+int(row['read_count'])
    if len(depth)<2 or min(depth.values())<=0:raise AcquisitionError('RNA pilot requires two positive-depth discovery brains')
    ordered=sorted(depth,key=lambda key:(depth[key],key))
    chosen={ordered[0],ordered[-1]}
    runs={row['run_id'] for row in rows if row['biological_id'] in chosen}
    selected=[a['asset_id'] for a in manifest['assets'] if a['source_id'] in ['ant_reference','mantis_rna','fly_reference'] or (a['source_id']=='ant_rna' and a.get('eligible') and a.get('run_accession') in runs)]
    if not selected:raise AcquisitionError('No eligible fixed RNA pilot assets')
    return selected


def download_manifest(manifest_path: str | Path, root: str | Path, *, sources: list[str] | None = None, max_bytes: int | None = None, max_files: int | None = None, bandwidth_bytes_per_second: int = 64 * 1024**2, max_seconds: float | None = None, receipt_path: str | Path | None = None,asset_ids: list[str] | None = None) -> dict:
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    validate_manifest(manifest)
    selected = [item for item in manifest["assets"] if (sources is None or item["source_id"] in sources) and (item.get("kind") != "fastq" or item.get("eligible"))]
    if asset_ids is not None:
        if len(asset_ids)!=len(set(asset_ids)) or set(asset_ids)-{a['asset_id'] for a in selected}:raise AcquisitionError('Invalid or ineligible explicit asset selection')
        selected=[a for a in selected if a['asset_id'] in asset_ids]
    if not selected:
        raise AcquisitionError("No eligible assets selected")
    if max_files is not None and len(selected) > max_files:
        raise AcquisitionError(f"Selected {len(selected)} files exceeds finite transfer limit {max_files}; split sources into bounded invocations")
    total = sum(int(item.get("bytes") or 0) for item in selected)
    if max_bytes is not None and (any(item.get("bytes") is None for item in selected) or total > max_bytes):
        raise AcquisitionError("Declared transfer size is unknown or exceeds the explicit byte ceiling")
    root = Path(root).resolve()
    selection_hash = canonical_hash([item["asset_id"] for item in selected])
    if receipt_path is None:
        receipt_path = root / "metadata" / f"download_receipts-{manifest['manifest_hash'][:16]}-{selection_hash[:12]}.json"
    else:
        receipt_path = Path(receipt_path)
        if not receipt_path.is_absolute():
            receipt_path = safe_destination(root, str(receipt_path))
        if not receipt_path.resolve().is_relative_to(root):
            raise AcquisitionError("Download receipt destination must remain in the study root")
    session = http_session()
    started = time.monotonic()
    report = {"manifest_hash": manifest["manifest_hash"], "selection_hash": selection_hash, "sources": sorted({item["source_id"] for item in selected}), "receipt_path": str(receipt_path), "files": []}
    for item in selected:
        remaining = max_seconds - (time.monotonic()-started) if max_seconds is not None else None
        if remaining is not None and remaining <= 0:
            raise AcquisitionError("Finite transfer runtime reached before the next file")
        report["files"].append(download_asset(item, root, session=session, bandwidth_bytes_per_second=bandwidth_bytes_per_second, max_seconds=remaining))
        atomic_json(receipt_path, report)
    return report


def safe_extract(archive: str | Path, destination: str | Path, *, max_unpacked_bytes: int = 100 * 1024**3, max_members: int = 100_000, include_pattern: str | None = None, materialize_internal_links: bool = False) -> list[dict]:
    """Extract ordinary files only; links, traversal, device files and bombs fail."""
    archive, destination = Path(archive), Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    included = re.compile(include_pattern) if include_pattern else None
    written: list[dict] = []
    total = 0
    if zipfile.is_zipfile(archive):
        with zipfile.ZipFile(archive) as handle:
            members = handle.infolist()
            if len(members) > max_members:
                raise AcquisitionError("Archive exceeds member limit")
            for member in members:
                safe_destination(destination, member.filename)
                mode = member.external_attr >> 16
                if stat.S_ISLNK(mode) or (stat.S_IFMT(mode) and not (stat.S_ISREG(mode) or stat.S_ISDIR(mode))):
                    raise AcquisitionError(f"Unsafe ZIP member type: {member.filename}")
                if not member.is_dir() and (included is None or included.search(member.filename)):
                    total += member.file_size
            if total > max_unpacked_bytes:
                raise AcquisitionError("Archive exceeds unpacked-size ceiling")
            for member in members:
                if member.is_dir() or (included is not None and not included.search(member.filename)):
                    continue
                target = safe_destination(destination, member.filename)
                target.parent.mkdir(parents=True, exist_ok=True)
                temporary = target.with_name(target.name + ".extracting")
                with handle.open(member) as source, temporary.open("wb") as out:
                    shutil.copyfileobj(source, out, length=8 * 1024 * 1024)
                if temporary.stat().st_size != member.file_size:
                    raise AcquisitionError("ZIP extracted size mismatch")
                temporary.replace(target)
                written.append({"path": str(target), "bytes": member.file_size, "sha256": file_hash(target)})
    elif tarfile.is_tarfile(archive):
        with tarfile.open(archive, "r:*") as handle:
            members = handle.getmembers()
            if len(members) > max_members:
                raise AcquisitionError("Archive exceeds member limit")
            index={posixpath.normpath(member.name):member for member in members}
            if len(index)!=len(members):raise AcquisitionError('Duplicate TAR member paths')
            def resolved_payload(member,seen=None):
                if member.isfile():return member
                if not materialize_internal_links or not (member.issym() or member.islnk()):raise AcquisitionError(f'Unsafe TAR member type: {member.name}')
                seen=set() if seen is None else seen
                if member.name in seen:raise AcquisitionError('Cyclic TAR internal link')
                seen.add(member.name)
                link=member.linkname
                if not link or PurePosixPath(link).is_absolute() or '\\' in link or ':' in link:raise AcquisitionError('Unsafe TAR internal link target')
                name=posixpath.normpath(posixpath.join(posixpath.dirname(member.name),link) if member.issym() else link)
                safe_destination(destination,name)
                if name not in index:raise AcquisitionError('TAR internal link target is absent from the archive')
                return resolved_payload(index[name],seen)
            payloads={}
            for member in members:
                safe_destination(destination, member.name)
                if member.isdir():continue
                payloads[member.name]=resolved_payload(member)
                if included is None or included.search(member.name):total+=payloads[member.name].size
            if total > max_unpacked_bytes:
                raise AcquisitionError("Archive exceeds unpacked-size ceiling")
            for member in members:
                if member.isdir() or (included is not None and not included.search(member.name)):
                    continue
                payload=payloads[member.name]
                target = safe_destination(destination, member.name)
                target.parent.mkdir(parents=True, exist_ok=True)
                temporary = target.with_name(target.name + ".extracting")
                with handle.extractfile(payload) as source, temporary.open("wb") as out:
                    shutil.copyfileobj(source, out, length=8 * 1024 * 1024)
                if temporary.stat().st_size != payload.size:
                    raise AcquisitionError("TAR extracted size mismatch")
                temporary.replace(target)
                written.append({"path": str(target), "bytes": payload.size, "sha256": file_hash(target),**({'materialized_internal_link':member.linkname} if not member.isfile() else {})})
    else:
        raise AcquisitionError(f"Unsupported archive: {archive}")
    return written


def inventory_file(path: str | Path) -> dict:
    path = Path(path)
    result: dict[str, Any] = {"path": str(path), "bytes": path.stat().st_size, "suffix": path.suffix.lower()}
    if path.suffix.lower() == ".mat":
        import h5py
        if h5py.is_hdf5(path):
            arrays = []
            with h5py.File(path, "r") as handle:
                def visit(name: str, obj: Any) -> None:
                    if isinstance(obj, h5py.Dataset):
                        arrays.append({"name": name, "shape": list(obj.shape), "dtype": str(obj.dtype), "matlab_class": str(obj.attrs.get("MATLAB_class", ""))})
                handle.visititems(visit)
            result.update(format="MATLAB-v7.3/HDF5", arrays=arrays)
        else:
            from scipy.io import whosmat
            result.update(format="MATLAB", arrays=[{"name": name, "shape": list(shape), "class": kind} for name, shape, kind in whosmat(path)])
    elif path.suffix.lower() in {".h5", ".hdf5"}:
        import h5py
        with h5py.File(path, "r") as handle:
            arrays = []
            handle.visititems(lambda name, obj: arrays.append({"name": name, "shape": list(obj.shape), "dtype": str(obj.dtype)}) if isinstance(obj, h5py.Dataset) else None)
        result.update(format="HDF5", arrays=arrays)
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        result["sheets"] = [{"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "first_rows": [[str(v) if v is not None else None for v in row] for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 5), values_only=True)]} for sheet in workbook]
        workbook.close()
    elif path.name.endswith((".csv", ".tsv", ".txt", ".csv.gz", ".tsv.gz", ".txt.gz")):
        opener = gzip.open if path.suffix == ".gz" else open
        with opener(path, "rt", encoding="utf-8-sig", errors="replace") as handle:
            result["first_lines"] = [handle.readline().rstrip("\n") for _ in range(5)]
    return result


def inventory(root: str | Path, output: str | Path) -> dict:
    files = [inventory_file(path) for path in sorted(Path(root).rglob("*")) if path.is_file() and not path.name.endswith((".part", ".tmp"))]
    report = {"created_utc": utc_now(), "root": str(Path(root).resolve()), "files": files, "status": "schema_inventory_only"}
    atomic_json(output, report)
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    discover = sub.add_parser("discover")
    discover.add_argument("--output", required=True)
    discover.add_argument("--sources", nargs="+")
    download = sub.add_parser("download")
    download.add_argument("--manifest", required=True)
    download.add_argument("--root", required=True)
    download.add_argument("--sources", nargs="+")
    download.add_argument("--max-bytes", type=int)
    download.add_argument("--max-files", type=int)
    download.add_argument("--bandwidth-mib", type=float, default=64)
    download.add_argument("--max-seconds", type=float)
    download.add_argument("--receipt-path")
    extract = sub.add_parser("extract")
    extract.add_argument("--archive", required=True)
    extract.add_argument("--destination", required=True)
    extract.add_argument("--max-unpacked-bytes", type=int, default=100 * 1024**3)
    extract.add_argument("--include-pattern")
    scan = sub.add_parser("inventory")
    scan.add_argument("--root", required=True)
    scan.add_argument("--output", required=True)
    args = parser.parse_args(argv)
    if args.command == "discover":
        result = discover_sources(args.output, args.sources)
        print(json.dumps({"manifest_hash": result["manifest_hash"], "assets": len(result["assets"])}))
    elif args.command == "download":
        result = download_manifest(args.manifest, args.root, sources=args.sources, max_bytes=args.max_bytes, max_files=args.max_files, bandwidth_bytes_per_second=int(args.bandwidth_mib*1024**2), max_seconds=args.max_seconds, receipt_path=args.receipt_path)
        print(json.dumps({"verified_files": len(result["files"])}))
    elif args.command == "extract":
        result = safe_extract(args.archive, args.destination, max_unpacked_bytes=args.max_unpacked_bytes, include_pattern=args.include_pattern)
        print(json.dumps({"extracted_files": len(result)}))
    else:
        result = inventory(args.root, args.output)
        print(json.dumps({"inventoried_files": len(result["files"])}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
