"""Native read-only VTL schema/clock inventory, without coordinate outcomes.

Published writer/order scripts are hashed as documentation, never imported or
executed. Position tokens are checked only for finite pair syntax; their values,
changes, repetition, minima, distances and behavioral summaries are not emitted.
Ordered slots have no validated biological identity or observable imputation mask.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime,timezone
import getpass
import hashlib
import json
import math
import os
from pathlib import Path
import re
import stat
import time
import zlib

ROOT=Path('/scratch/pwa209/horsehair-worm')
ACQUISITION=ROOT/'maintenance/study-af-vtl-whole-member-canary-v1-20261021/acquisition-result.json'
ACQUISITION_SHA='0c91c8a596402cc3ee7dc7abc4e8fe38d51e3453654fb3224b24e63b6d771808'
ACQUISITION_HELPER_SHA='2a866bb7842111b7fdcb7a56e05cf77e8ce6a866277f4b808d7f761a6f7a8ae1'
INDEX=ROOT/'maintenance/study-af-vtl-directory-dataset01-v1-20261021/metadata-index.json'
INDEX_SHA='58209aca20832ec3897e30014393ace5dd25b186d4580c05d14921bb2f92a9e7'
MAPPING=ROOT/'maintenance/study-af-additional-public-metadata-v2-20261021/vtl_information.xlsx'
MAPPING_SHA='b53bcc2242a6019f6556d217116c98d021878b6d4ed88c281f752d1eb555518d'
WRITER_SHA='d83e7431c258977ba07e78b9f7430fe3e8a5dfb5b7569a5cfb9740e643f8d114'
ORDER_SHA='ff011406d399640da05b82c37256ea9d4a6f69f8b48e8c1ef93927ce43a3fe64'
TARGET='2019-11-10_2_11-order.txt'
RAW_BYTES=30121971
RAW_SHA='6587355772ae4d730bf9839ef2c1b60b9dbddbe852a4a7f24e611e39b0336996'
CRC='abefabb9'
SOURCE_BYTES={'index':1016200,'mapping':571760,'writer':7622,'order':3173}
URL='https://vtl.biocuckoo.cn/static/data/dataset1.zip'
ETAG='"14746eac0-5f88f9022eaf8"'
MAX_RAW=64*1024**2
MAX_LINE=65536
MAX_ROWS=2_000_000
MAX_FIELDS=1024
FALSE_FLAGS=('scientific_result','analysis_admitted','admission_authority','effects_inspected',
    'coordinate_values_exported','coordinate_numeric_values_aggregated','provider_code_executed',
    'physical_2d_units_qualified','absolute_or_exposure_clock_qualified','biological_identity_qualified',
    'per_slot_observation_mask_recoverable','ordered_track_imputation_qualified','data_repaired_or_imputed',
    'source_replication_passed','biological_transfer_supported','h1_h2_h3_gates_cleared',
    'representative_cohort_selected','gene_control_mapping_verified')


def require(value,message):
    if not value:raise ValueError(message)


def hashes(path):
    h=hashlib.sha256();crc=length=0
    with Path(path).open('rb') as stream:
        for block in iter(lambda:stream.read(65536),b''):
            length+=len(block);require(length<=MAX_RAW,'Finite whole-file byte bound exceeded')
            h.update(block);crc=zlib.crc32(block,crc)
    return {'sha256':h.hexdigest(),'bytes':length,'crc32':f'{crc:08x}'}


def ordinary(path,*,fresh=False):
    p=Path(path);root=ROOT.resolve(strict=True);resolved=p.resolve()
    prefix=ROOT if p.is_relative_to(ROOT) else root
    require(p.is_absolute() and p.is_relative_to(prefix) and resolved!=root and resolved.is_relative_to(root)
        and all(not x.is_symlink() for x in [p,*p.parents] if x!=prefix and x.is_relative_to(prefix)),
        'Exact ordinary personal path required')
    if fresh:require(not p.exists(),'Fresh output required; no replay')
    else:
        s=p.stat();require(stat.S_ISREG(s.st_mode) and s.st_nlink==1,'Ordinary single-link source required')
    return resolved


def read_json(path,max_bytes=2_000_000):
    require(Path(path).stat().st_size<=max_bytes,'Bounded source JSON required')
    def pairs(items):
        result={}
        for key,value in items:
            require(key not in result,'Duplicate JSON member refused');result[key]=value
        return result
    return json.loads(Path(path).read_text(),object_pairs_hook=pairs,
        parse_constant=lambda value:(_ for _ in ()).throw(ValueError('Nonfinite JSON refused')))


def exclusive(path,value):
    raw=(json.dumps(value,sort_keys=True,indent=2,allow_nan=False)+'\n').encode()
    require(len(raw)<=2_000_000,'Bounded structural result required')
    with Path(path).open('xb') as stream:stream.write(raw);stream.flush();os.fsync(stream.fileno())


def native_context():
    job=os.environ.get('SLURM_JOB_ID','')
    require(re.fullmatch('[1-9][0-9]*',job) and getpass.getuser()=='pwa209','Owned Slurm audit allocation required')
    require(os.environ.get('SLURM_CPUS_PER_TASK')=='1','Exactly1CPU required')
    for key in ('SLURM_JOB_GPUS','SLURM_STEP_GPUS','CUDA_VISIBLE_DEVICES'):
        require(os.environ.get(key,'') in {'','-1','NoDevFiles','none','None'},'Any GPU assignment including device0 refused')
    end=float(os.environ.get('SLURM_JOB_END_TIME','nan'))
    require(math.isfinite(end) and time.time()+240<=end,'Native finalization margin required')
    return {'job_id':job,'cpus':1,'gpus':0,'allocation_end_unix':end}


def code_documentation(writer,order):
    writer_text=Path(writer).read_text(encoding='utf-8');order_text=Path(order).read_text(encoding='utf-8')
    facts=[(writer_text,'fps = cap.get(cv2.CAP_PROP_FPS)'),
        (writer_text,'fly.write(str(cou)'),(writer_text,'(430,430)'),
        (writer_text,"str(i-dic1[fn][4])"),(order_text,'w.write(l0)'),
        (order_text,'if sp3[i0]==0:sp3[i0]=xy0')]
    require(all(needle in text for text,needle in facts),'Exact reviewed code-documentation structures required')
    line=lambda text,needle:next(i for i,s in enumerate(text.splitlines(),1) if needle in s)
    return {'writer':{'url':'https://raw.githubusercontent.com/BioCUCKOO/VTL/main/1.vtl.py','sha256':WRITER_SHA,
            'nominal_video_fps_code_line':line(writer_text,facts[0][1]),
            'resized_crop_code_line':line(writer_text,'(430,430)')},
        'ordering':{'url':'https://raw.githubusercontent.com/BioCUCKOO/VTL/main/2.dataorder.py','sha256':ORDER_SHA,
            'header_copy_code_line':line(order_text,'w.write(l0)'),
            'unmatched_slot_carry_forward_code_line':line(order_text,'if sp3[i0]==0:sp3[i0]=xy0')},
        'code_executed':False,'code_version_that_produced_this_member_verified':False,
        'documented_writer_shape':'7header fields; coordinate pairs followed by frame-index and detection-count fields',
        'documented_coordinates':'resized crop pixels in published writer; physical units of this member unqualified',
        'documented_ordering':'fixed slots initialized on first row; unmatched slots reuse previous coordinates',
        'explicit_per_slot_observed_or_imputed_flag_documented':False}


def source_inputs(acquisition,index,mapping,writer,order):
    require(Path(acquisition)==ACQUISITION and Path(index)==INDEX and Path(mapping)==MAPPING,
        'Exact native canary acquisition/directory/workbook paths required')
    pins=[]
    for role,path,sha,size in [('acquisition',acquisition,ACQUISITION_SHA,None),
        ('index',index,INDEX_SHA,SOURCE_BYTES['index']),('mapping',mapping,MAPPING_SHA,SOURCE_BYTES['mapping']),
        ('writer',writer,WRITER_SHA,SOURCE_BYTES['writer']),('order',order,ORDER_SHA,SOURCE_BYTES['order'])]:
        p=ordinary(path);measured=hashes(p)
        require(measured['sha256']==sha and (size is None or measured['bytes']==size),'Exact documentation/source bytes required: '+role)
        pins.append({'role':role,'path':str(p),'sha256':sha,'bytes':measured['bytes']})
    receipt=read_json(acquisition)
    require(receipt.get('schema')=='vtl_single_whole_member_acquisition_v1'
        and receipt.get('schema_version')==1 and receipt.get('status')=='complete_whole_member_crc_verified_first_sha256_recorded'
        and receipt.get('helper_sha256')==ACQUISITION_HELPER_SHA and receipt.get('final_promoted') is True
        and receipt.get('url')==URL and receipt.get('archive_bytes')==5490797248 and receipt.get('etag')==ETAG,
        'Exact positive complete whole-member acquisition receipt required')
    for key in ('scientific_result','admission_authority','full_archive_verified','cryptographic_provider_checksum_verified',
        'raw_track_values_parsed','representative_cohort_selected','gene_control_mapping_verified',
        'individual_identity_clock_units_verified','ordered_track_imputation_qualified','data_repaired_or_imputed',
        'source_replication_passed','biological_transfer_supported'):
        require(receipt.get(key) is False,'Acquisition must not clear scientific/identity gates')
    member=ordinary(receipt.get('member',''))
    require(member==ordinary(ACQUISITION.parent/TARGET) and member.name==TARGET,'Exact frozen whole member required')
    verification=receipt.get('verification',{})
    require(verification=={'compressed_bytes':5326466,'raw_bytes':RAW_BYTES,'crc32':CRC,'sha256':RAW_SHA,
        'sha256_basis':'first_complete_acquisition_sha256_not_provider_checksum','deposited_zip_crc_verified':True,
        'cryptographic_provider_member_checksum_available':False,'cryptographic_provider_checksum_verified':False},
        'First-acquisition SHA/deposited CRC must retain their actual provenance')
    directory=read_json(index)
    selected=[e for e in directory.get('members',[]) if e.get('name')==TARGET]
    require(len(selected)==1 and selected[0]==receipt.get('selected_member')
        and selected[0].get('raw_bytes')==RAW_BYTES and selected[0].get('crc32')==CRC,
        'Exact selected central record must match verified directory')
    require(receipt.get('source_pins')==[{'path':str(index),'bytes':SOURCE_BYTES['index'],'sha256':INDEX_SHA},
        {'path':str(mapping),'bytes':SOURCE_BYTES['mapping'],'sha256':MAPPING_SHA}], 'Original exact directory/workbook bindings required')
    selection=receipt.get('selection',{})
    require(selection.get('rule')=='minimum_sha256_of_ascii_filename_among_frozen_bounded_supported_members_v1'
        and selection.get('selected_filename_sha256')==hashlib.sha256(TARGET.encode('ascii')).hexdigest()
        and selection.get('eligible_members')==777 and selection.get('eligible_filename_roster_sha256')==
            'ad00db36e8c7317d3a2683125c67d6a70516552fa95e2fe478844e14a1dcfaec'
        and selection.get('outcome_values_used') is False and selection.get('representative_cohort_claim') is False,
        'Actual outcome-blind bounded parser-probe selection required')
    measured=hashes(member)
    require(measured=={'sha256':RAW_SHA,'bytes':RAW_BYTES,'crc32':CRC},'Whole raw member changed since first acquisition')
    return member,measured,pins,code_documentation(writer,order)


def finite_number(text):
    if len(text)>64:return None
    try:value=float(text)
    except (ValueError,OverflowError):return None
    return value if math.isfinite(value) else None


def header_shape(fields):
    if len(fields)!=7:return None
    total,fps=finite_number(fields[0]),finite_number(fields[1])
    pair=lambda s:bool(re.fullmatch('[0-9]{1,7}:[0-9]{1,7}',s))
    if (total is None or not 0<total<=1e9 or fps is None or not 0<fps<=1000
        or not pair(fields[2]) or not pair(fields[3])
        or not all(re.fullmatch('[+-]?[0-9]{1,8}',s) for s in fields[4:])):return None
    return {'declared_total_video_frames_literal':fields[0],'declared_video_fps_literal':fields[1],
        'published_writer_crop_field_metadata':[{'field_index_1based':i+1,
            'syntax':'unsigned_integer_pair_colon','sha256':hashlib.sha256(fields[i].encode()).hexdigest(),
            'literal_exported':False} for i in (2,3)],
        'published_writer_remaining_instrument_literals':fields[4:],
        'field_semantics_from_matching_published_code_not_verified_production_history':True,
        'video_fps_is_an_absolute_or_exposure_clock':False,'physical_units_qualified':False}


def mapping_matches(mapping,member_name):
    import openpyxl
    book=openpyxl.load_workbook(mapping,read_only=True,data_only=False,keep_links=False)
    try:
        require(len(book.sheetnames)==1,'One original filename-mapping sheet required')
        sheet=book[book.sheetnames[0]]
        require(sheet.max_row<=30000 and sheet.max_column<=10,'Bounded original mapping dimensions required')
        rows=sheet.iter_rows();header=next(rows)
        require([c.value for c in header]==['th_number','gene_id','vlt_name'],'Original literal filename/gene/line columns required')
        targets={member_name:'exact_ordered_filename',member_name.removesuffix('-order.txt')+'.txt':'published_order_writer_raw_filename_relation',
            member_name.removesuffix('-order.txt'):'literal_stem_match_only'}
        matches=[]
        for row in rows:
            first=row[0]
            if first.data_type=='f' or type(first.value) is not str or first.value not in targets:continue
            values=[]
            for column,cell in enumerate(row,1):
                coordinate=f'{openpyxl.utils.get_column_letter(column)}{first.row}'
                if cell.data_type=='f':values.append({'cell':coordinate,'type':'formula_not_evaluated','literal':None})
                elif cell.value is None or type(cell.value) in (str,int,float):
                    require(cell.value is None or len(str(cell.value))<=256,'Bounded matched metadata literals required')
                    if type(cell.value) is float:require(math.isfinite(cell.value),'Finite source metadata literal required')
                    values.append({'cell':coordinate,'type':'source_literal','literal':cell.value})
                else:values.append({'cell':coordinate,'type':'unsupported_source_type','literal':None})
            matches.append({'match_kind':targets[first.value],'source_sheet':sheet.title,'source_row':first.row,'cells':values})
            require(len(matches)<=10,'Bounded exact filename metadata matches required')
        return {'candidate_filename_relations':targets,'literal_matches':matches,'matched_rows':len(matches),
            'control_edges_inferred':False,'gene_identifiers_validated':False,'biological_unit_inferred':False}
    finally:book.close()


def inspect_track(path,deadline):
    counters=Counter();slots=Counter();spacing=Counter();first=last=previous=None;min_clock=max_clock=None
    first_slots=None;recognized=None;read_bytes=0;header=None;spacing_other=0
    with Path(path).open('rb') as stream:
        for number in range(MAX_ROWS+2):
            require(time.monotonic()<deadline,'Finite structural parsing deadline reached')
            raw=stream.readline(MAX_LINE+1)
            if not raw:break
            read_bytes+=len(raw);require(read_bytes<=MAX_RAW and len(raw)<=MAX_LINE,'Finite file/line bounds exceeded')
            require(number<=MAX_ROWS,'Finite row bound exceeded')
            if number==0:
                try:fields=raw.decode('utf-8').rstrip('\r\n').split('\t')
                except UnicodeDecodeError:fields=[]
                header=header_shape(fields);recognized=header is not None
                counters['header_fields']=len(fields);continue
            counters['data_rows']+=1
            try:fields=raw.decode('utf-8').rstrip('\r\n').split('\t')
            except UnicodeDecodeError:
                counters['undecodable_rows']+=1;previous=None;continue
            if not recognized:
                counters['uninterpreted_rows']+=1;continue
            if not 2<=len(fields)<=MAX_FIELDS:
                counters['unsupported_field_count_rows']+=1;previous=None;continue
            prefix=fields[:-2];count=len(prefix)
            slots[count]+=1
            if number==1:first_slots=count
            valid_pairs=empty_pairs=invalid_pairs=0
            for token in prefix:
                if token=='':empty_pairs+=1;continue
                parts=token.split(',')
                if len(parts)==2 and all(finite_number(p) is not None for p in parts):valid_pairs+=1
                else:invalid_pairs+=1
            counters['finite_coordinate_pair_tokens']+=valid_pairs;counters['empty_coordinate_tokens']+=empty_pairs
            counters['malformed_or_nonfinite_coordinate_tokens']+=invalid_pairs
            if invalid_pairs:counters['rows_with_malformed_or_nonfinite_coordinate_tokens']+=1
            if empty_pairs:counters['rows_with_empty_coordinate_tokens']+=1
            if not re.fullmatch('[0-9]{1,12}',fields[-2]):
                counters['invalid_frame_index_rows']+=1;previous=None
            else:
                clock=int(fields[-2]);counters['valid_frame_index_rows']+=1
                if first is None:first=clock
                last=clock;min_clock=clock if min_clock is None else min(min_clock,clock);max_clock=clock if max_clock is None else max(max_clock,clock)
                if previous is not None:
                    delta=clock-previous
                    if delta==0:counters['adjacent_duplicate_frame_indices']+=1
                    if delta<0:counters['adjacent_frame_index_reversals']+=1
                    if delta>1:counters['adjacent_positive_frame_gaps_gt1']+=1
                    if delta in spacing or len(spacing)<64:spacing[delta]+=1
                    else:spacing_other+=1
                previous=clock
            if not re.fullmatch('[0-9]{1,6}',fields[-1]):counters['invalid_detection_count_rows']+=1
            else:
                detection=int(fields[-1]);counters['valid_detection_count_rows']+=1
                if detection==0:counters['zero_detection_count_rows']+=1
                counters['detector_count_'+('below' if detection<count else 'equals' if detection==count else 'above')+'_coordinate_slots_rows']+=1
        else:raise ValueError('Finite row bound exceeded')
    return {'status':'recognized_writer_shape_inventory' if recognized else 'unknown_header_schema_inventory',
        'header_shape_matches_published_writer':bool(recognized),'header_instrument_metadata':header,
        'bytes_read':read_bytes,'counts':dict(sorted(counters.items())),
        'coordinate_slot_count_histogram':{str(k):v for k,v in sorted(slots.items())},
        'first_row_coordinate_slot_count':first_slots,
        'frame_index_metadata':{'first':first,'last':last,'minimum':min_clock,'maximum':max_clock,
            'adjacent_spacing_histogram_first64_distinct':{str(k):v for k,v in sorted(spacing.items())},
            'spacing_values_outside_histogram_count':spacing_other,'expected_grid_missingness_qualified':False,
            'seconds_or_exposure_times_derived':False},
        'position_tokens_examined_for_structure_only':bool(recognized),
        'missing_observation_interpretation':'empty/malformed syntax and original detector-count field only; exact per-slot missingness and carry-forward events are unobservable',
        'slot_roster_is_validated_individual_identity':False,'unchanged_coordinates_interpreted_as_observed_or_still':False,
        'all_slots_observed_when_detection_count_equals_slots':False,**{k:False for k in FALSE_FLAGS}}


def audit(acquisition,index,mapping,writer,order,output,*,max_seconds=900):
    require(type(max_seconds) is int and 60<=max_seconds<=1800,'Finite60..1800-second structural audit required')
    native=native_context();started=time.monotonic();deadline=min(started+max_seconds,started+native['allocation_end_unix']-time.time()-180)
    own_sha=hashes(Path(__file__))['sha256']
    member,measured,pins,documentation=source_inputs(*map(Path,(acquisition,index,mapping,writer,order)))
    output=ordinary(output,fresh=True);require((ROOT/'maintenance').resolve() in output.parents,'Fresh personal maintenance case required')
    output.mkdir()
    base={'schema_version':1,'schema':'vtl_read_only_track_structure_v1','native':native,'auditor_sha256':own_sha,
        'source_pins':pins,'acquisition':{'path':str(acquisition),'sha256':ACQUISITION_SHA},
        'member':{'path':str(member),**measured},'published_code_documentation':documentation,
        'cryptographic_provider_member_checksum_available':False,'first_acquisition_sha256_not_provider_checksum':True,
        **{k:False for k in FALSE_FLAGS}}
    exclusive(output/'audit-start.json',base)
    try:
        inventory=inspect_track(member,deadline);inventory['filename_metadata']=mapping_matches(Path(mapping),member.name)
        require(source_inputs(*map(Path,(acquisition,index,mapping,writer,order)))==(member,measured,pins,documentation)
            and hashes(Path(__file__))['sha256']==own_sha,'Member/documentation/auditor changed during inventory')
        require(time.monotonic()<deadline,'Finite audit finalization deadline reached')
        exclusive(output/'schema-and-clock-inventory.json',inventory)
        result={**base,'status':'complete_bounded_outcome_blind_vtl_schema_inventory',
            'inventory':{'path':str(output/'schema-and-clock-inventory.json'),'sha256':hashes(output/'schema-and-clock-inventory.json')['sha256']},
            'all_source_bytes_reverified':True,'elapsed_seconds':time.monotonic()-started,
            'unknown_schema_does_not_clear_scientific_gates':True}
        exclusive(output/'receipt.json',result);return result
    except BaseException as error:
        exclusive(output/'failure.json',{**base,'status':'incomplete_structural_inventory','error_type':type(error).__name__,
            'error':str(error),'elapsed_seconds':time.monotonic()-started})
        raise


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    for key in ('acquisition','index','mapping','writer','order','output'):p.add_argument('--'+key,type=Path,required=True)
    p.add_argument('--max-seconds',type=int,default=900);a=p.parse_args()
    result=audit(a.acquisition,a.index,a.mapping,a.writer,a.order,a.output,max_seconds=a.max_seconds)
    print(json.dumps({'status':result['status'],'inventory':result['inventory'],'scientific_result':False}))
