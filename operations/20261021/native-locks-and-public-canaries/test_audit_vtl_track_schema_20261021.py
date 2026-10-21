"""Synthetic-only VTL structural inventory tests; no provider requests/data."""
import ast
import hashlib
import importlib.util
import json
from pathlib import Path
import time
import zlib

import openpyxl
import pytest

SOURCE=Path(__file__).with_name('audit-vtl-track-schema-20261021.py')
spec=importlib.util.spec_from_file_location('vtl_schema_private',SOURCE)
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)

HEADER='1000.0\t20.0\t0:480\t0:480\t0\t55\t12\n'
TRACK=HEADER+'10,11\t20,21\t2\t2\n10,11\t\t3\t1\nnan,1\tNaN,2\t5\t0\n12,13\t30,31\t5\t2\n1,2\t3,4\t4\t2\n'


@pytest.fixture
def fixture(tmp_path,monkeypatch):
    root=tmp_path/'personal';case=root/'maintenance/acquired';case.mkdir(parents=True)
    folder=root/'maintenance/docs';folder.mkdir()
    member=case/m.TARGET;member.write_bytes(TRACK.encode())
    mapping=folder/'vtl_information.xlsx';book=openpyxl.Workbook();sheet=book.active;sheet.title='Mapping'
    sheet.append(['th_number','gene_id','vlt_name'])
    sheet.append([m.TARGET.removesuffix('-order.txt'),'SYNTHETIC_GENE','SYNTHETIC_LINE'])
    sheet.append(['unrelated','not_selected','SECRET_OUTCOME_UNRELATED'])
    book.save(mapping);book.close()
    writer=folder/'VTL-1.vtl.py';writer.write_text("# synthetic documentation, never executed\nfps = cap.get(cv2.CAP_PROP_FPS)\nfly.write(str(cou))\n# (430,430)\n# str(i-dic1[fn][4])\nraise RuntimeError('NEVER_EXECUTE')\n")
    order=folder/'VTL-2.dataorder.py';order.write_text("# synthetic documentation, never executed\nw.write(l0)\nif sp3[i0]==0:sp3[i0]=xy0\nraise RuntimeError('NEVER_EXECUTE')\n")
    raw=m.hashes(member);selected={'name':m.TARGET,'raw_bytes':raw['bytes'],'crc32':raw['crc32'],'flags':0,'method':8}
    index=folder/'metadata-index.json';index.write_text(json.dumps({'members':[selected]}))
    receipt={'schema':'vtl_single_whole_member_acquisition_v1','schema_version':1,
        'status':'complete_whole_member_crc_verified_first_sha256_recorded','helper_sha256':m.ACQUISITION_HELPER_SHA,
        'final_promoted':True,'url':m.URL,'archive_bytes':5490797248,'etag':m.ETAG,'member':str(member),
        'verification':{'compressed_bytes':5326466,'raw_bytes':raw['bytes'],'crc32':raw['crc32'],'sha256':raw['sha256'],
            'sha256_basis':'first_complete_acquisition_sha256_not_provider_checksum','deposited_zip_crc_verified':True,
            'cryptographic_provider_member_checksum_available':False,'cryptographic_provider_checksum_verified':False},
        'selected_member':selected,'selection':{'rule':'minimum_sha256_of_ascii_filename_among_frozen_bounded_supported_members_v1',
            'selected_filename_sha256':hashlib.sha256(m.TARGET.encode()).hexdigest(),'eligible_members':777,
            'eligible_filename_roster_sha256':'ad00db36e8c7317d3a2683125c67d6a70516552fa95e2fe478844e14a1dcfaec',
            'outcome_values_used':False,'representative_cohort_claim':False},
        'source_pins':[{'path':str(index),'bytes':index.stat().st_size,'sha256':m.hashes(index)['sha256']},
            {'path':str(mapping),'bytes':mapping.stat().st_size,'sha256':m.hashes(mapping)['sha256']}]}
    receipt.update({key:False for key in ('scientific_result','admission_authority','full_archive_verified',
        'cryptographic_provider_checksum_verified','raw_track_values_parsed','representative_cohort_selected',
        'gene_control_mapping_verified','individual_identity_clock_units_verified','ordered_track_imputation_qualified',
        'data_repaired_or_imputed','source_replication_passed','biological_transfer_supported')})
    acquisition=case/'acquisition-result.json';acquisition.write_text(json.dumps(receipt))
    for name,path in [('ACQUISITION',acquisition),('INDEX',index),('MAPPING',mapping)]:monkeypatch.setattr(m,name,path)
    for name,path in [('ACQUISITION_SHA',acquisition),('INDEX_SHA',index),('MAPPING_SHA',mapping),('WRITER_SHA',writer),('ORDER_SHA',order)]:
        monkeypatch.setattr(m,name,m.hashes(path)['sha256'])
    monkeypatch.setattr(m,'ROOT',root);monkeypatch.setattr(m,'RAW_BYTES',raw['bytes']);monkeypatch.setattr(m,'RAW_SHA',raw['sha256']);monkeypatch.setattr(m,'CRC',raw['crc32'])
    monkeypatch.setattr(m,'SOURCE_BYTES',{role:path.stat().st_size for role,path in [('index',index),('mapping',mapping),('writer',writer),('order',order)]})
    monkeypatch.setattr(m,'native_context',lambda:{'job_id':'123','cpus':1,'gpus':0,'allocation_end_unix':time.time()+3600})
    return {'root':root,'member':member,'acquisition':acquisition,'index':index,'mapping':mapping,'writer':writer,'order':order,
        'receipt':receipt,'output':root/'maintenance/audit'}


def invoke(f):return m.audit(f['acquisition'],f['index'],f['mapping'],f['writer'],f['order'],f['output'])


def test_recognized_schema_quality_and_clock_counts_without_coordinate_values(fixture):
    inventory=m.inspect_track(fixture['member'],time.monotonic()+30)
    counts=inventory['counts']
    assert inventory['status']=='recognized_writer_shape_inventory' and counts['data_rows']==5
    assert counts['finite_coordinate_pair_tokens']==7 and counts['empty_coordinate_tokens']==1
    assert counts['malformed_or_nonfinite_coordinate_tokens']==2
    assert counts['adjacent_duplicate_frame_indices']==1 and counts['adjacent_frame_index_reversals']==1
    assert counts['adjacent_positive_frame_gaps_gt1']==1 and counts['zero_detection_count_rows']==1
    assert inventory['frame_index_metadata']['first']==2 and inventory['frame_index_metadata']['last']==4
    assert inventory['frame_index_metadata']['adjacent_spacing_histogram_first64_distinct']=={'-1':1,'0':1,'1':1,'2':1}
    assert inventory['coordinate_slot_count_histogram']=={'2':5} and inventory['first_row_coordinate_slot_count']==2
    raw=json.dumps(inventory)
    assert all(token not in raw for token in ['10,11','20,21','12,13','30,31','0:480'])
    assert all(x['literal_exported'] is False for x in inventory['header_instrument_metadata']['published_writer_crop_field_metadata'])
    assert all(inventory[k] is False for k in m.FALSE_FLAGS)
    assert inventory['all_slots_observed_when_detection_count_equals_slots'] is False


def test_source_bound_full_audit_copies_no_raw_values_and_executes_no_code(fixture):
    f=fixture;before={key:m.hashes(f[key]) for key in ('member','acquisition','index','mapping','writer','order')}
    result=invoke(f)
    assert result['status']=='complete_bounded_outcome_blind_vtl_schema_inventory'
    assert result['all_source_bytes_reverified'] is True and all(result[k] is False for k in m.FALSE_FLAGS)
    inventory=json.loads(Path(result['inventory']['path']).read_text())
    assert inventory['filename_metadata']['matched_rows']==1
    cells=inventory['filename_metadata']['literal_matches'][0]['cells']
    assert [x['literal'] for x in cells]==[m.TARGET.removesuffix('-order.txt'),'SYNTHETIC_GENE','SYNTHETIC_LINE']
    assert 'SECRET_OUTCOME_UNRELATED' not in json.dumps(result)+json.dumps(inventory)
    assert result['published_code_documentation']['ordering']['unmatched_slot_carry_forward_code_line']==3
    assert result['published_code_documentation']['code_version_that_produced_this_member_verified'] is False
    assert before=={key:m.hashes(f[key]) for key in before}
    assert not (f['output']/'failure.json').exists()


@pytest.mark.parametrize('header',['SECRET_OUTCOME\t99\n','1\tNaN\t0:480\t0:480\t0\t55\t12\n',
    '1\t1e309\t0:480\t0:480\t0\t55\t12\n','1\t20\tunknown\t0:480\t0\t55\t12\n'])
def test_unknown_header_remains_uninterpreted_without_echoing_values(tmp_path,header):
    path=tmp_path/'synthetic.txt';path.write_text(header+'SECRET_ROW_A\tSENSITIVE_NUMBER_123\n')
    inventory=m.inspect_track(path,time.monotonic()+30)
    assert inventory['status']=='unknown_header_schema_inventory' and inventory['header_instrument_metadata'] is None
    assert inventory['counts']['uninterpreted_rows']==1
    assert inventory['frame_index_metadata']['first'] is None
    assert all(s not in json.dumps(inventory) for s in ('SECRET_OUTCOME','SECRET_ROW_A','SENSITIVE_NUMBER_123'))
    assert inventory['position_tokens_examined_for_structure_only'] is False


def test_unknown_binary_header_is_inventory_only_not_unsafe_decode(tmp_path):
    path=tmp_path/'synthetic.txt';path.write_bytes(b'\xff\xfeunknown\n\xff\xfe\n')
    inventory=m.inspect_track(path,time.monotonic()+30)
    assert inventory['status']=='unknown_header_schema_inventory' and inventory['counts']['undecodable_rows']==1


@pytest.mark.parametrize('kind',['invalid_frame','invalid_count','variable_slots','empty_fields','bad_pair'])
def test_malformed_missing_and_variable_rows_retained_without_repair(tmp_path,kind):
    rows={'invalid_frame':'1,2\tclock_unknown\t2\n','invalid_count':'1,2\t2\tunknown\n',
        'variable_slots':'1,2\t3,4\t5,6\t2\t3\n','empty_fields':'\t2\t0\n','bad_pair':'SECRET_COORDINATE\t2\t1\n'}
    path=tmp_path/'synthetic.txt';path.write_text(HEADER+rows[kind])
    inventory=m.inspect_track(path,time.monotonic()+30)
    assert inventory['counts']['data_rows']==1 and 'SECRET_COORDINATE' not in json.dumps(inventory)
    assert all(inventory[k] is False for k in m.FALSE_FLAGS)


def test_invalid_clock_row_does_not_create_false_adjacent_spacing_across_it(tmp_path):
    path=tmp_path/'synthetic.txt';path.write_text(HEADER+'1,2\t1\t1\n1,2\tunknown\t1\n1,2\t9\t1\n')
    inventory=m.inspect_track(path,time.monotonic()+30)
    assert inventory['frame_index_metadata']['adjacent_spacing_histogram_first64_distinct']=={}
    assert inventory['counts']['invalid_frame_index_rows']==1


def test_invalid_first_row_does_not_invent_initial_slot_roster(tmp_path):
    path=tmp_path/'synthetic.txt';path.write_text(HEADER+'unknown\n1,2\t9\t1\n')
    inventory=m.inspect_track(path,time.monotonic()+30)
    assert inventory['first_row_coordinate_slot_count'] is None


@pytest.mark.parametrize('kind',['line','bytes','rows','deadline'])
def test_all_parser_resource_bounds_are_finite(tmp_path,monkeypatch,kind):
    path=tmp_path/'synthetic.txt';path.write_text(TRACK)
    if kind=='line':monkeypatch.setattr(m,'MAX_LINE',8)
    if kind=='bytes':monkeypatch.setattr(m,'MAX_RAW',8)
    if kind=='rows':monkeypatch.setattr(m,'MAX_ROWS',1)
    with pytest.raises(ValueError):m.inspect_track(path,time.monotonic()+(-1 if kind=='deadline' else 30))


@pytest.mark.parametrize('key',['member','acquisition','index','mapping','writer','order'])
def test_changed_source_bytes_prevent_completed_audit(fixture,key):
    f=fixture;f[key].write_bytes(f[key].read_bytes()+b'changed')
    with pytest.raises(ValueError):invoke(f)
    assert not (f['output']/'receipt.json').exists()


@pytest.mark.parametrize('key',['status','helper','provider_hash_claim','scientific_gate','selected','source_pins','selection'])
def test_dishonest_acquisition_provenance_refused(fixture,monkeypatch,key):
    f=fixture;r=f['receipt']
    if key=='status':r['status']='partial'
    if key=='helper':r['helper_sha256']='a'*64
    if key=='provider_hash_claim':r['verification']['sha256_basis']='provider_checksum'
    if key=='scientific_gate':r['individual_identity_clock_units_verified']=True
    if key=='selected':r['selected_member']['name']='different.txt'
    if key=='source_pins':r['source_pins'][0]['sha256']='a'*64
    if key=='selection':r['selection']['outcome_values_used']=True
    f['acquisition'].write_text(json.dumps(r));monkeypatch.setattr(m,'ACQUISITION_SHA',m.hashes(f['acquisition'])['sha256'])
    with pytest.raises(ValueError):invoke(f)


def test_mid_inspection_source_change_preserves_start_and_failure(fixture,monkeypatch):
    f=fixture;original=m.inspect_track
    def changed(path,deadline):
        result=original(path,deadline);path.write_bytes(path.read_bytes()+b'changed');return result
    monkeypatch.setattr(m,'inspect_track',changed)
    with pytest.raises(ValueError):invoke(f)
    assert (f['output']/'audit-start.json').exists() and (f['output']/'failure.json').exists()
    assert not (f['output']/'receipt.json').exists()


def test_matched_metadata_formulas_not_executed_or_exported_as_values(fixture):
    f=fixture;book=openpyxl.load_workbook(f['mapping']);book.active['B2']='=1+1';book.save(f['mapping']);book.close()
    result=m.mapping_matches(f['mapping'],m.TARGET)
    assert result['literal_matches'][0]['cells'][1]=={'cell':'B2','type':'formula_not_evaluated','literal':None}
    assert result['gene_identifiers_validated'] is False and result['control_edges_inferred'] is False


def test_mapping_only_literal_name_relations_preserves_duplicates_no_gene_guess(fixture):
    f=fixture;book=openpyxl.load_workbook(f['mapping']);sheet=book.active
    sheet.append([m.TARGET,'DIFFERENT_GENE',None]);sheet.append([m.TARGET.removesuffix('-order.txt')+'.txt',None,'another_line'])
    book.save(f['mapping']);book.close()
    result=m.mapping_matches(f['mapping'],m.TARGET)
    assert result['matched_rows']==3
    assert {r['match_kind'] for r in result['literal_matches']}=={'exact_ordered_filename','literal_stem_match_only','published_order_writer_raw_filename_relation'}
    assert result['biological_unit_inferred'] is False


def test_existing_output_refused_without_overwrite(fixture):
    f=fixture;f['output'].mkdir();(f['output']/'sentinel').write_bytes(b'preserve')
    with pytest.raises(ValueError):invoke(f)
    assert (f['output']/'sentinel').read_bytes()==b'preserve'


@pytest.mark.parametrize('gpu',['0','1','0,1','gpu:1'])
def test_native_contract_refuses_all_gpu_assignments(fixture,monkeypatch,gpu):
    original=importlib.util.spec_from_file_location('native_contract_check',SOURCE);module=importlib.util.module_from_spec(original);original.loader.exec_module(module)
    monkeypatch.setenv('SLURM_JOB_ID','123');monkeypatch.setenv('SLURM_CPUS_PER_TASK','1');monkeypatch.setenv('SLURM_JOB_END_TIME',str(time.time()+3600))
    monkeypatch.setattr(module.getpass,'getuser',lambda:'pwa209');monkeypatch.setenv('CUDA_VISIBLE_DEVICES',gpu)
    with pytest.raises(ValueError,match='GPU'):module.native_context()


def test_source_is_readonly_and_never_imports_or_executes_provider_code():
    tree=ast.parse(SOURCE.read_text())
    assert not any(isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id in {'eval','exec','compile','__import__'} for n in ast.walk(tree))
    imports={a.name for n in ast.walk(tree) if isinstance(n,ast.Import) for a in n.names}
    assert not imports & {'numpy','pandas','torch','cv2','matplotlib','runpy','importlib','subprocess'}
